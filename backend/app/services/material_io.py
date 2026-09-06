"""Material catalog import/export (JSON / CSV / XLSX).

Shared write path for batch import: parse → normalize → dedupe
(cas > smiles > norm_key) → dry-run preview or commit via MaterialStore.
"""
from __future__ import annotations

import csv
import io
import json
import uuid
from dataclasses import dataclass, field
from typing import Any, Iterable

from loguru import logger

from ..db.material_store import get_material_store, norm_key
from ..domain.knowledge import RAW_MATERIALS

EXPORT_COLUMNS = (
    "name",
    "role",
    "zh_name",
    "cas_no",
    "smiles",
    "formula",
    "molar_mass",
    "price_cny_per_kg",
    "voc_contrib",
    "density_gcm3",
    "functional_class",
    "substitute_group",
    "supplier",
    "availability",
    "origin",
)

_FLOAT_FIELDS = {
    "molar_mass",
    "price_cny_per_kg",
    "voc_contrib",
    "density_gcm3",
    "oil_absorption",
    "tg_k",
    "equivalent_weight",
    "hansen_d",
    "hansen_p",
    "hansen_h",
    "hlb",
}
_INT_FIELDS = {"lead_time_days"}
_BOOL_FIELDS = {"svhc", "water_compatible", "archived"}
_AVAILABILITY = {"in_stock", "restricted", "discontinued"}


@dataclass
class ImportRowResult:
    name: str
    action: str  # create | update | skip | error
    reason: str = ""
    matched_by: str = ""
    existing_name: str = ""


@dataclass
class ImportPreview:
    batch_id: str
    total: int = 0
    creates: int = 0
    updates: int = 0
    skips: int = 0
    errors: int = 0
    rows: list[ImportRowResult] = field(default_factory=list)


def _coerce_value(key: str, value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        if value == "":
            return None
    if key in _FLOAT_FIELDS:
        try:
            return float(value)
        except (TypeError, ValueError):
            return None
    if key in _INT_FIELDS:
        try:
            return int(float(value))
        except (TypeError, ValueError):
            return None
    if key in _BOOL_FIELDS:
        if isinstance(value, bool):
            return value
        s = str(value).strip().lower()
        if s in {"1", "true", "yes", "y", "是"}:
            return True
        if s in {"0", "false", "no", "n", "否"}:
            return False
        return None
    return value


def normalize_record(raw: dict[str, Any]) -> dict[str, Any] | None:
    if not isinstance(raw, dict):
        return None
    aliases = {
        "material": "name",
        "material_name": "name",
        "名称": "name",
        "cas": "cas_no",
        "CAS": "cas_no",
        "CAS号": "cas_no",
        "中文名": "zh_name",
        "角色": "role",
        "供应商": "supplier",
    }
    data: dict[str, Any] = {}
    for k, v in raw.items():
        key = aliases.get(str(k).strip(), str(k).strip())
        if key == "name" or key in EXPORT_COLUMNS or key in _FLOAT_FIELDS or key in _BOOL_FIELDS:
            data[key] = v
    name = str(data.get("name") or "").strip()
    if not name:
        return None
    out: dict[str, Any] = {"name": name}
    for key, value in data.items():
        if key == "name":
            continue
        coerced = _coerce_value(key, value)
        if coerced is None:
            continue
        if key == "availability" and str(coerced) not in _AVAILABILITY:
            continue
        out[key] = coerced
    out.setdefault("role", "")
    out.setdefault("availability", "in_stock")
    return out


def parse_json_bytes(payload: bytes) -> list[dict[str, Any]]:
    data = json.loads(payload.decode("utf-8-sig"))
    if isinstance(data, dict) and "materials" in data:
        data = data["materials"]
    if not isinstance(data, list):
        raise ValueError('JSON 根节点须为数组，或 {"materials": [...]}')
    return [r for r in data if isinstance(r, dict)]


def parse_csv_bytes(payload: bytes) -> list[dict[str, Any]]:
    text = payload.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    return [dict(row) for row in reader]


def parse_xlsx_bytes(payload: bytes) -> list[dict[str, Any]]:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise ValueError("服务器未安装 openpyxl，无法解析 xlsx") from exc
    wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        return []
    keys = [str(h).strip() if h is not None else "" for h in header]
    out: list[dict[str, Any]] = []
    for row in rows_iter:
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        item = {keys[i]: row[i] for i in range(min(len(keys), len(row))) if keys[i]}
        out.append(item)
    return out


def parse_xls_bytes(payload: bytes) -> list[dict[str, Any]]:
    try:
        import xlrd  # type: ignore
    except ImportError as exc:
        raise ValueError("暂不支持直接解析 .xls，请另存为 .xlsx 或 .csv 后再导入") from exc
    book = xlrd.open_workbook(file_contents=payload)
    sheet = book.sheet_by_index(0)
    if sheet.nrows < 1:
        return []
    keys = [str(sheet.cell_value(0, c)).strip() for c in range(sheet.ncols)]
    out: list[dict[str, Any]] = []
    for r in range(1, sheet.nrows):
        item = {keys[c]: sheet.cell_value(r, c) for c in range(sheet.ncols) if keys[c]}
        out.append(item)
    return out


def detect_and_parse(filename: str, payload: bytes) -> list[dict[str, Any]]:
    name = (filename or "").lower()
    if name.endswith(".json"):
        return parse_json_bytes(payload)
    if name.endswith(".csv"):
        return parse_csv_bytes(payload)
    if name.endswith(".xlsx"):
        return parse_xlsx_bytes(payload)
    if name.endswith(".xls"):
        return parse_xls_bytes(payload)
    head = payload.lstrip()[:1]
    if head in (b"{", b"["):
        return parse_json_bytes(payload)
    if b"," in payload[:200] or b";" in payload[:200]:
        return parse_csv_bytes(payload)
    raise ValueError("无法识别格式，请使用 json / csv / xlsx（xls 请另存为 xlsx）")


def _find_existing(store, record: dict[str, Any]) -> tuple[str | None, str]:
    cas = str(record.get("cas_no") or "").strip()
    if cas:
        hit = store.find_by_cas(cas)
        if hit is not None:
            return hit.name, "cas_no"
    smiles = str(record.get("smiles") or "").strip()
    if smiles:
        hit = store.find_by_smiles(smiles)
        if hit is not None:
            return hit.name, "smiles"
    key = norm_key(record["name"])
    for name, spec in RAW_MATERIALS.items():
        if norm_key(name) == key:
            return name, "norm_key"
        if cas and str(spec.get("cas_no") or "").strip() == cas:
            return name, "cas_no"
        if smiles and str(spec.get("smiles") or "").strip() == smiles:
            return name, "smiles"
    hit = store.get(record["name"])
    if hit is not None:
        return hit.name, "norm_key"
    return None, ""


def plan_import(records: Iterable[dict[str, Any]], *, batch_id: str | None = None) -> ImportPreview:
    store = get_material_store()
    preview = ImportPreview(batch_id=batch_id or str(uuid.uuid4()))
    for raw in records:
        preview.total += 1
        record = normalize_record(raw)
        if record is None:
            preview.errors += 1
            preview.rows.append(ImportRowResult(name="", action="error", reason="缺少 name"))
            continue
        existing, matched_by = _find_existing(store, record)
        if existing:
            preview.updates += 1
            preview.rows.append(
                ImportRowResult(
                    name=record["name"],
                    action="update",
                    matched_by=matched_by,
                    existing_name=existing,
                )
            )
        else:
            preview.creates += 1
            preview.rows.append(
                ImportRowResult(name=record["name"], action="create", matched_by="")
            )
    return preview


def commit_import(
    records: Iterable[dict[str, Any]],
    *,
    batch_id: str | None = None,
    origin: str = "import",
) -> ImportPreview:
    store = get_material_store()
    preview = ImportPreview(batch_id=batch_id or str(uuid.uuid4()))
    meta = {"import_batch_id": preview.batch_id}
    for raw in records:
        preview.total += 1
        record = normalize_record(raw)
        if record is None:
            preview.errors += 1
            preview.rows.append(ImportRowResult(name="", action="error", reason="缺少 name"))
            continue
        name = record.pop("name")
        existing, matched_by = _find_existing(store, {**record, "name": name})
        target_name = existing or name
        try:
            regulatory = (
                dict(record.get("regulatory") or {})
                if isinstance(record.get("regulatory"), dict)
                else {}
            )
            regulatory.update(meta)
            record["regulatory"] = regulatory
            # Re-import always restores visibility.
            record["archived"] = False
            ok = store.upsert(target_name, record, origin=origin, overwrite=True)
            if not ok:
                preview.errors += 1
                preview.rows.append(
                    ImportRowResult(name=name, action="error", reason="写入失败")
                )
                continue
            if existing:
                preview.updates += 1
                action = "update"
            else:
                preview.creates += 1
                action = "create"
            preview.rows.append(
                ImportRowResult(
                    name=name,
                    action=action,
                    matched_by=matched_by,
                    existing_name=existing or "",
                )
            )
        except Exception as exc:
            logger.warning("material import row failed: {} ({})", name, exc)
            preview.errors += 1
            preview.rows.append(
                ImportRowResult(name=name, action="error", reason=str(exc)[:200])
            )
    RAW_MATERIALS.refresh()
    return preview


def export_records(
    *,
    q: str = "",
    role: str = "",
    availability: str = "",
    functional_class: str = "",
    substitute_group: str = "",
    include_archived: bool = False,
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    term = q.strip().lower()
    for name, spec in RAW_MATERIALS.items():
        if not include_archived and spec.get("archived"):
            continue
        if role and spec.get("role") != role:
            continue
        if availability and (spec.get("availability") or "in_stock") != availability:
            continue
        if functional_class and (spec.get("functional_class") or "") != functional_class:
            continue
        if substitute_group and (spec.get("substitute_group") or "") != substitute_group:
            continue
        if term and term not in name.lower() and term not in str(spec.get("zh_name") or "").lower():
            continue
        item: dict[str, Any] = {"name": name}
        for col in EXPORT_COLUMNS:
            if col == "name":
                continue
            if col == "origin":
                item[col] = spec.get("origin") or "seed"
                continue
            if col in spec and spec[col] is not None:
                item[col] = spec[col]
        rows.append(item)
    return rows


def serialize_json(records: list[dict[str, Any]]) -> bytes:
    return json.dumps({"materials": records}, ensure_ascii=False, indent=2).encode("utf-8")


def serialize_csv(records: list[dict[str, Any]]) -> bytes:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=list(EXPORT_COLUMNS), extrasaction="ignore")
    writer.writeheader()
    for row in records:
        writer.writerow({k: row.get(k, "") for k in EXPORT_COLUMNS})
    return buf.getvalue().encode("utf-8-sig")


def serialize_xlsx(records: list[dict[str, Any]]) -> bytes:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise ValueError("服务器未安装 openpyxl，无法导出 xlsx") from exc
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "materials"
    ws.append(list(EXPORT_COLUMNS))
    for row in records:
        ws.append([row.get(c, "") for c in EXPORT_COLUMNS])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def import_template(fmt: str) -> tuple[bytes, str, str]:
    fmt = fmt.lower().strip()
    if fmt == "json":
        return serialize_json([]), "application/json", "materials_template.json"
    if fmt == "csv":
        return serialize_csv([]), "text/csv", "materials_template.csv"
    if fmt in {"xlsx", "xls"}:
        return (
            serialize_xlsx([]),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "materials_template.xlsx",
        )
    raise ValueError("format 须为 json / csv / xlsx")
