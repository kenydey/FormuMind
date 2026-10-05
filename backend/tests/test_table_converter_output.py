"""Tables as the converters actually hand them over (round-4).

The table tests all feed ``table_contract`` hand-written markdown with a clean header row. What the parsers
produce is not that:

* **.docx / .pptx** — Word tables carry no header markup, so MarkItDown renders them with an *empty* header
  row and pushes the real one into the body (``|  |  |  |`` / ``| 项目 | 指标 | 单位 |``). The asset got
  ``headers == ['', '', '']``; the header row became a bogus property, the unit column could not be found
  (**every value lost its unit**: ``65`` instead of ``65 %``), and the kind came from the caption alone.
* **PDF** (pymupdf4llm) bolds header cells (``**项目**``) and breaks wrapped cells with ``<br>``; a caption
  promoted to a heading arrives as ``## 表1 典型性能``. All of it ended up in the stored text.
* A datasheet table whose caption and headers name no topic (``项目 | 指标 | 单位`` on a spreadsheet,
  ``Property | Value | Unit`` in an English datasheet) was ``other`` — the normalizer, which finds columns by
  exactly those words, never saw it.
* ``耐盐雾性能`` (the usual wording of the salt-spray row) was the one row of a Chinese TDS left unmapped.

The last section parses real files and skips when the parser extras are not installed.
"""
from __future__ import annotations

import io

import pytest

from app.services import parsing, table_contract as tc, table_normalize as tn

DOCX_STYLE = """# 环氧锌磷酸盐底漆 技术数据表（TDS）

表1 典型性能

|  |  |  |
| --- | --- | --- |
| 项目 | 指标 | 单位 |
| 固体含量 | 65 | % |
| 粘度 | 1200 | mPa·s |
| 密度 | 1.35 | g/cm3 |
| 耐盐雾性能 | 720 | h |
"""


def _extract(md: str, parser: str = "markitdown"):
    return tc.extract_tables("s1", tc.blocks_from_markdown(md), parser=parser)


def _props(assets):
    sets = tn.normalize_tables(assets)
    assert len(sets) == 1
    return {p.name: p for p in sets[0].properties}, sets[0].warnings


# ── header row ──────────────────────────────────────────────────────────────

def test_a_blank_header_row_gives_way_to_the_first_body_row():
    (asset,) = _extract(DOCX_STYLE)
    assert asset.headers == ["项目", "指标", "单位"]
    assert asset.rows[0] == ["固体含量", "65", "%"]
    assert len(asset.rows) == 4


def test_the_recovered_header_lets_the_normalizer_keep_every_unit():
    (asset,) = _extract(DOCX_STYLE)
    assert asset.kind == "performance"
    props, warnings = _props([asset])
    assert "项目" not in props, "the header row must not become a property"
    assert {n: (p.value, p.unit_normalized) for n, p in props.items()} == {
        "固体含量": (65.0, "%"),
        "粘度": (1200.0, "mPa.s"),
        "密度": (1.35, "g/cm3"),
        "耐盐雾性能": (720.0, "h"),
    }
    assert warnings == []


def test_a_headerless_data_table_keeps_its_first_row_as_data():
    md = "|  |  |  |\n| --- | --- | --- |\n| 固体含量 | 65 | % |\n| 粘度 | 1200 | mPa·s |\n"
    (asset,) = _extract(md)
    assert asset.headers == ["", "", ""]
    assert asset.rows[0] == ["固体含量", "65", "%"]


def test_a_textual_key_value_table_is_not_mistaken_for_a_header():
    md = "|  |  |\n| --- | --- |\n| 外观 | 灰色液体 |\n| 状态 | 双组分 |\n| 颜色 | 灰色 |\n"
    (asset,) = _extract(md)
    assert asset.headers == ["", ""]
    assert len(asset.rows) == 3


def test_a_two_row_table_is_left_alone():
    md = "|  |  |\n| --- | --- |\n| 项目 | 指标 |\n"
    (asset,) = _extract(md)
    assert asset.headers == ["", ""] and asset.rows == [["项目", "指标"]]


def test_a_real_header_is_never_replaced():
    md = "| 项目 | 指标 |\n| --- | --- |\n| 项目 | 指标 |\n| 固体含量 | 65 |\n"
    (asset,) = _extract(md)
    assert asset.headers == ["项目", "指标"]
    assert asset.rows[0] == ["项目", "指标"]


# ── cell and caption text ───────────────────────────────────────────────────

def test_emphasis_and_line_breaks_do_not_reach_the_stored_cells():
    md = "|**项目**|**指标<br>(h)**|__单位__|\n|---|---|---|\n|**耐盐雾性能**|720|h|\n"
    (asset,) = _extract(md, parser="hybrid")
    assert asset.headers == ["项目", "指标 (h)", "单位"]
    assert asset.rows == [["耐盐雾性能", "720", "h"]]
    assert "**" in asset.raw_markdown, "the untouched text stays available"


def test_emphasis_inside_a_cell_is_left_alone():
    assert tc._clean_cell("a **b** c") == "a **b** c"
    assert tc._clean_cell("**b**") == "b"
    assert tc._clean_cell("x<br/>y") == "x y"


def test_a_caption_loses_its_heading_marker():
    md = "## 表1 典型性能\n\n| 项目 | 指标 |\n| --- | --- |\n| 粘度 | 1200 |\n"
    (asset,) = _extract(md)
    assert asset.caption == "表1 典型性能"


# ── classification without topical keywords ─────────────────────────────────

@pytest.mark.parametrize(
    "headers",
    [
        ["项目", "指标", "单位"],
        ["检验项目", "技术要求"],
        ["Property", "Value", "Unit"],
        ["Item", "Specification"],
    ],
)
def test_a_name_column_next_to_a_value_column_is_a_property_table(headers):
    kind, conf = tc.classify_table(headers)
    assert kind == "performance"
    assert 0.5 <= conf < 0.65, "inferred from layout, so below a two-keyword match"


@pytest.mark.parametrize(
    "headers",
    [["A", "B"], ["标准"], ["项目"], ["Value", "Unit"], ["指标名称"], ["spherical particles"]],
)
def test_other_tables_are_still_other(headers):
    assert tc.classify_table(headers)[0] == "other"


def test_the_normalizer_and_the_classifier_share_one_vocabulary():
    assert tn._NAME_COL_HINTS is tc.NAME_COL_HINTS
    assert tn._VALUE_COL_HINTS is tc.VALUE_COL_HINTS
    assert tn._UNIT_COL_HINTS is tc.UNIT_COL_HINTS


@pytest.mark.parametrize(
    "name",
    ["耐盐雾性能", "耐盐雾性", "耐盐雾试验", "耐中性盐雾", "盐雾性能", "Salt spray resistance", "Neutral salt spray (h)", "NSS"],
)
def test_the_salt_spray_row_maps_however_a_datasheet_words_it(name):
    assert tn.normalize_property_name(name) == "salt_spray"


# ── real files (skipped without the parser extras) ──────────────────────────

ROWS = [("项目", "指标", "单位"), ("固体含量", "65", "%"), ("粘度", "1200", "mPa·s"), ("耐盐雾性能", "720", "h")]


def _assert_datasheet(result: parsing.ParseResult):
    assert len(result.tables) == 1, result.markdown
    (asset,) = result.tables
    assert asset.headers == ["项目", "指标", "单位"], asset.headers
    assert asset.caption == "表1 典型性能", asset.caption
    props, warnings = _props([asset])
    assert {n: (p.value, p.unit_normalized) for n, p in props.items()} == {
        "固体含量": (65.0, "%"),
        "粘度": (1200.0, "mPa.s"),
        "耐盐雾性能": (720.0, "h"),
    }
    assert warnings == []


def test_a_real_docx_datasheet_keeps_its_header_and_its_units():
    pytest.importorskip("docx")
    pytest.importorskip("markitdown")
    import docx

    doc = docx.Document()
    doc.add_heading("环氧锌磷酸盐底漆 技术数据表（TDS）", 1)
    doc.add_paragraph("表1 典型性能")
    table = doc.add_table(rows=len(ROWS), cols=3)
    for i, row in enumerate(ROWS):
        for j, value in enumerate(row):
            table.cell(i, j).text = value
    buf = io.BytesIO()
    doc.save(buf)
    result = parsing.parse_document(buf.getvalue(), "docx")
    assert result.parser == "markitdown"
    _assert_datasheet(result)


def test_a_real_pptx_datasheet_has_a_clean_caption():
    pytest.importorskip("pptx")
    pytest.importorskip("markitdown")
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "表1 典型性能"
    table = slide.shapes.add_table(len(ROWS), 3, Inches(1), Inches(2), Inches(6), Inches(2)).table
    for i, row in enumerate(ROWS):
        for j, value in enumerate(row):
            table.cell(i, j).text = value
    buf = io.BytesIO()
    prs.save(buf)
    _assert_datasheet(parsing.parse_document(buf.getvalue(), "pptx"))


def test_a_real_pdf_datasheet_has_clean_headers_and_caption():
    fitz = pytest.importorskip("fitz")
    pytest.importorskip("pymupdf4llm")
    from pathlib import Path

    font = next(
        (p for p in ("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc", "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc") if Path(p).is_file()),
        None,
    )
    if font is None:
        pytest.skip("no CJK font on this machine to draw the Chinese fixture")
    doc = fitz.open()
    page = doc.new_page()
    page.insert_font(fontname="cjk", fontfile=font)
    page.insert_text((60, 70), "环氧锌磷酸盐底漆 技术数据表（TDS）", fontname="cjk", fontsize=16)
    page.insert_text((60, 105), "表1 典型性能", fontname="cjk", fontsize=11)
    x0, y0, cw, rh = 60, 120, 120, 24
    for i, row in enumerate(ROWS):
        for j, value in enumerate(row):
            rect = fitz.Rect(x0 + j * cw, y0 + i * rh, x0 + (j + 1) * cw, y0 + (i + 1) * rh)
            page.draw_rect(rect, color=(0, 0, 0), width=0.8)
            page.insert_text((rect.x0 + 6, rect.y0 + 16), value, fontname="cjk", fontsize=10.5)
    result = parsing.parse_document(doc.tobytes(), "pdf")
    assert result.parser == "hybrid"
    _assert_datasheet(result)


# ── a column-per-property table is not a property list ──────────────────────

WIDE_DOCX = """表2 性能指标（合并表头）

|  |  |  |  |
| --- | --- | --- | --- |
| 物理性能 | | 防腐性能 | |
| 固含量(%) | 粘度(mPa·s) | 盐雾(h) | 附着力(级) |
| 65 | 1200 | 720 | 1 |
| 60 | 1100 | 500 | 2 |
"""


def test_a_wide_table_is_skipped_instead_of_producing_properties_named_after_numbers():
    sets = tn.normalize_tables(_extract(WIDE_DOCX))
    assert [(s.properties, s.warnings) for s in sets] == [
        ([], ["skipped: column-per-property (wide) table — its row labels are numbers"])
    ]


def test_a_property_list_with_an_occasional_numeric_label_is_still_normalised():
    md = "| 项目 | 指标 | 单位 |\n| --- | --- | --- |\n| 固体含量 | 65 | % |\n| 粘度 | 1200 | mPa·s |\n| 3 | 4 | h |\n"
    (asset,) = _extract(md)
    props, _ = _props([asset])
    assert {"固体含量", "粘度"} <= set(props)
